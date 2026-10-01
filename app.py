from flask import Flask, request, jsonify, send_from_directory
from pptx import Presentation
from pptx.util import Pt
from werkzeug.utils import secure_filename
import os
import uuid

app = Flask(__name__)

OUTPUT_FOLDER = os.path.join(os.getcwd(), "generated_ppts")
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


@app.route("/", methods=["GET"])
def home():
    return jsonify({
        "status": "success",
        "message": "PPT Generator API is running"
    })


@app.route("/createppt", methods=["POST"])
def create_ppt():

    data = request.get_json(silent=True)

    if not data:
        return jsonify({
            "status": "error",
            "message": "Request body must contain valid JSON"
        }), 400

    presentation_title = data.get(
        "presentationTitle",
        data.get("title", "Generated Presentation")
    )

    slides_data = data.get("slides", [])

    if not slides_data:
        slides_data = [
            {
                "title": presentation_title,
                "content": [
                    "Presentation generated through Copilot Studio",
                    "Custom connector successfully invoked",
                    "PowerPoint file created by the custom API"
                ]
            }
        ]

    presentation = Presentation()

    # Title slide
    title_slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    title_slide.shapes.title.text = presentation_title

    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = (
            "Created by Custom Tools Test V3"
        )

    # Content slides
    for slide_data in slides_data:

        slide_title = slide_data.get("title", "Untitled Slide")
        slide_content = slide_data.get("content", [])

        if isinstance(slide_content, str):
            slide_content = [slide_content]

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = slide_title

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        for index, bullet in enumerate(slide_content):

            if index == 0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()

            paragraph.text = str(bullet)
            paragraph.font.size = Pt(20)
            paragraph.level = 0

    safe_title = secure_filename(presentation_title)

    if not safe_title:
        safe_title = "Generated_Presentation"

    file_name = f"{safe_title}_{uuid.uuid4().hex[:8]}.pptx"
    file_path = os.path.join(OUTPUT_FOLDER, file_name)

    presentation.save(file_path)

    base_url = request.url_root.rstrip("/")
    download_url = f"{base_url}/download/{file_name}"

    return jsonify({
        "status": "success",
        "message": "PowerPoint presentation created successfully",
        "fileName": file_name,
        "downloadUrl": download_url,
        "slideCount": len(presentation.slides)
    })


@app.route("/download/<path:file_name>", methods=["GET"])
def download_ppt(file_name):

    return send_from_directory(
        OUTPUT_FOLDER,
        file_name,
        as_attachment=True,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        )
    )


if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000))
    )
